#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Confere o TRADE contra o Power BI "final" — sem gravar nada.

Duas conferências, com os arquivos que alimentavam o Power BI (unidade Z:):

  1. VENDA de um mês: a exportação do WinThor em  Z:\\TRADE MARKETING\\Power BI\\Geral\\Final\\AAAA-MM.xlsx
     contra a mesma regra do painel aplicada direto no Oracle (PCMOV, operações S e ED,
     QT x (PUNIT + VLOUTROS), filiais 1, 7, 10, 11 e 12).
  2. RUPTURA: a planilha EEG.xlsx (foto do estoque do Power BI) com a regra do painel
     (ativo, não promocional, entrada nos últimos 365 dias, estoque total disponível zerado).

Uso:
    python conferir_power_bi.py              # mês 2026-02 (o último que bate com o arquivo do Power BI)
    python conferir_power_bi.py 2026-01

Resultado de 2026-10-14 (para comparar depois; valores absolutos ficam fora do Git, que é público):
    venda de fev/2026 ........ painel +0,005% contra o Power BI
    clientes com venda ....... +4 · SKUs −0,5% · pares −1,0%
    ruptura (EEG de 02/03) ... +4,1%
"""
from __future__ import annotations

import os
import sys
import warnings
from datetime import date, datetime
from pathlib import Path

warnings.filterwarnings("ignore")
import oracledb
import openpyxl
from dotenv import load_dotenv

AQUI = Path(__file__).resolve().parent
PASTA = Path(r"Z:\TRADE MARKETING\Power BI\Geral")
FILIAIS = ("1", "7", "10", "11", "12")


def carregar_env():
    cands = []
    if os.environ.get("LUBE_ENV"):
        cands.append(Path(os.environ["LUBE_ENV"]))
    cands += [AQUI / "ENV", AQUI.parent / "ENV", AQUI.parent / "1 - CONFIGURACAO" / "ENV"]
    for c in cands:
        if c.is_file():
            load_dotenv(c)
            return
    raise SystemExit("ENV não encontrado (veja ENV.example).")


def conectar():
    dsn = oracledb.makedsn(os.environ["ORACLE_HOST"], int(os.environ.get("ORACLE_PORT", "1521")), service_name=os.environ["ORACLE_SERVICE_NAME"])
    return oracledb.connect(user=os.environ["ORACLE_USER"], password=os.environ["ORACLE_PASSWORD"], dsn=dsn)


def pct(a: float, b: float) -> str:
    return f"{(a / b - 1) * 100:+.3f}%" if b else "—"


def conferir_venda(cur, ym: str) -> None:
    arq = PASTA / "Final" / f"{ym}.xlsx"
    if not arq.is_file():
        print(f"[venda] arquivo do Power BI não encontrado: {arq}")
        return
    print(f"[venda] lendo {arq.name} (leva uns 30 s)…")
    ws = openpyxl.load_workbook(arq, read_only=True, data_only=True)[openpyxl.load_workbook(arq, read_only=True).sheetnames[0]]
    total, clientes, clientes_pos, skus, pares = 0.0, set(), set(), set(), set()
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i == 0 or r[8] is None or not str(r[8]).strip().isdigit():
            continue                                    # só linhas de dado (Seq. numérico)
        venda = float(r[20] or 0)
        total += venda
        cli, prod = r[5], r[9]
        clientes.add(cli); skus.add(prod); pares.add((cli, prod))
        if venda > 0:
            clientes_pos.add(cli)
    ano, mes = int(ym[:4]), int(ym[5:])
    ini = date(ano, mes, 1)
    fim = date(ano + (mes == 12), mes % 12 + 1, 1)
    cur.execute(f"""
        SELECT SUM(DECODE(codoper,'S',1,-1)*qt*(punit+NVL(vloutros,0))), COUNT(DISTINCT CASE WHEN codoper='S' THEN codcli END),
               COUNT(DISTINCT CASE WHEN codoper='S' THEN codprod END), COUNT(DISTINCT CASE WHEN codoper='S' THEN codcli||'-'||codprod END)
          FROM pcmov WHERE codoper IN ('S','ED') AND dtcancel IS NULL AND codfilial IN ({", ".join("'" + f + "'" for f in FILIAIS)})
           AND dtmov >= :a AND dtmov < :b""", a=datetime.combine(ini, datetime.min.time()), b=datetime.combine(fim, datetime.min.time()))
    v, c, s, p = cur.fetchone()
    print(f"[venda] {ym}:  venda líquida  painel R$ {float(v):,.2f}   Power BI R$ {total:,.2f}   ({pct(float(v), total)})")
    print(f"        clientes com venda  painel {c}   Power BI {len(clientes_pos)} (com devolução: {len(clientes)})")
    print(f"        SKUs  painel {s}   Power BI {len(skus)}      pares cliente x produto  painel {p}   Power BI {len(pares)}")


def conferir_ruptura(cur) -> None:
    eeg = PASTA / "Consultas" / "EEG.xlsx"
    cls = PASTA / "Consultas" / "Categoria e subcategoria.xlsx"
    if not eeg.is_file() or not cls.is_file():
        print("[ruptura] EEG.xlsx ou a classificação não estão em Z:")
        return
    import pandas as pd
    e = pd.read_excel(eeg, sheet_name="Export")
    e.columns = [str(c).strip() for c in e.columns]
    e = e[e["CODPROD"].notna()].copy()
    c = pd.read_excel(cls, sheet_name="Planilha1")
    c.columns = ["cod", "ean", "desc", "dep", "cat", "sub", "marca", "promo"]
    c = c[c["cod"].notna()]
    promo = {int(k): v for k, v in zip(c["cod"], c["promo"].fillna(""))}
    ref = datetime.fromtimestamp(eeg.stat().st_mtime)
    cur.execute("SELECT codprod, MAX(dtmov) FROM pcmov WHERE codoper='E' AND dtcancel IS NULL AND dtmov <= :d GROUP BY codprod", d=ref)
    ult = {int(a): b for a, b in cur.fetchall()}
    total = e["EST_DISPONIVEL"].fillna(0) + e["EST BLOQ S/AVARIA"].fillna(0)
    ativo = e["STATUS"].eq("ATIVO")
    prom = e["CODPROD"].map(lambda x: promo.get(int(x), "") == "P")
    rec = e["CODPROD"].map(lambda x: int(x) in ult and (ref - ult[int(x)]).days < 365)
    base = ativo & ~prom
    print(f"[ruptura] EEG de {ref:%d/%m/%Y}: SKUs ativos não promocionais {int(base.sum())} (Power BI: 4.314) · em ruptura {int((base & rec & (total <= 0)).sum())} (Power BI: 490)")


def main() -> int:
    carregar_env()
    ym = sys.argv[1] if len(sys.argv) > 1 else "2026-02"
    conn = conectar()
    cur = conn.cursor()
    conferir_venda(cur, ym)
    conferir_ruptura(cur)
    conn.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
