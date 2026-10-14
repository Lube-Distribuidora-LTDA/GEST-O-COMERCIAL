#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
TRADE — Gestão Comercial · construção do pacote de dados do painel.

Lê o WinThor (Oracle, só leitura), a classificação do Trade Marketing e a pasta de fotos
(FOTOS_ION) e monta o PACOTE DE DADOS do painel:

    meta.json          carimbo da carga, período, contagens
    dims.json          dimensões (produtos, clientes, RCAs, fornecedores...)
    venda.bin.gz       fato de venda, mensal por RCA x cliente x produto (colunar)
    estoque.json       base de produtos: estoque, ruptura, sem venda, imagem
    entradas.json      custo de entrada por produto x mês
    cc.json            conta corrente de fornecedor (verba)
    historico.json     foto diária dos indicadores (cresce a cada carga)
    imagens.json       código do produto → arquivo na pasta FOTOS_ION

Este módulo é BIBLIOTECA: `construir()` devolve o pacote na memória. Quem grava no banco
(DATA WAREHOUSE, schema trade) é o `sync_trade.py`, via `gravar_trade.py`. Rodado direto,
serve só para gerar os arquivos em ../web/dados/ (modo sem banco, para teste):

    python extrair_trade.py              # carga completa em arquivos
    python extrair_trade.py --usar-cache # reaproveita a extração de venda (só Oracle pequeno)

Credenciais: arquivo ENV (mesmo formato dos outros BIs) na pasta do script, na
pasta acima, em "1 - CONFIGURACAO", ou apontado pela variável LUBE_ENV.
NUNCA grava senha em lugar nenhum.

Regras que o painel herdou do BI "final" do Power BI estão no README, com as
divergências conscientes (valor / valor_original).
"""
from __future__ import annotations

import gzip
import json
import os
import re
import struct
import sys
import time
import unicodedata
from array import array
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

import oracledb
from dotenv import load_dotenv

AQUI = Path(__file__).resolve().parent
SAIDA = (AQUI.parent / "web" / "dados")
CACHE = AQUI / "_cache_venda.bin"

# --------------------------------------------------------------------------
# Parâmetros de negócio (um lugar só)
# --------------------------------------------------------------------------
FILIAIS_VENDA = ("1", "7", "10", "11", "12")      # filiais que vendem ao cliente
FILIAIS_ESTOQUE = (1, 2, 8, 7, 10)                # mesmas do ESTOQUE x VENDA do COMPRAS
DATA_INICIAL = date(2023, 1, 1)                   # histórico de venda
DATA_INICIAL_CC = date(2024, 1, 1)                # lançamentos da conta corrente exibidos
COMPRADORES_EXCLUIDOS = (621, 278, 376, 375, 277, 178)   # igual ao BI COMPRAS
DIAS_VENDA_JANELA = 56                            # 8 semanas (ruptura / sem venda)
# Planilha do Trade (mantida à mão). Tarefa agendada não enxerga unidade mapeada (Z:), então o caminho de rede do mesmo
# arquivo vem do ENV (TRADE_CLASSIFICACAO) como segunda opção. Nada de endereço de servidor no código: o repositório é público.
ARQUIVOS_CLASSIFICACAO = (
    Path(r"Z:\TRADE MARKETING\Power BI\Geral\Consultas\Categoria e subcategoria.xlsx"),
)


def log(msg: str) -> None:
    # sob o sync_trade.py o logging já está configurado (arquivo + tela); rodando sozinho, só a tela
    import logging
    if logging.getLogger().handlers:
        logging.getLogger("bi").info(msg)
    else:
        print(f"[{datetime.now():%H:%M:%S}] {msg}", flush=True)


# --------------------------------------------------------------------------
# Conexão
# --------------------------------------------------------------------------
def carregar_env() -> None:
    candidatos = []
    if os.environ.get("LUBE_ENV"):
        candidatos.append(Path(os.environ["LUBE_ENV"]))
    for pasta in (AQUI, AQUI.parent, AQUI.parent / "1 - CONFIGURACAO"):
        for nome in ("ENV", ".env", "env"):
            candidatos.append(pasta / nome)
    for c in candidatos:
        if c.is_file():
            load_dotenv(c)
            log(f"configuração carregada de {c.name} ({c.parent.name})")
            return
    log("AVISO: nenhum ENV encontrado — usando só variáveis de ambiente.")


def conectar():
    dsn = oracledb.makedsn(os.environ["ORACLE_HOST"], int(os.environ.get("ORACLE_PORT", "1521")),
                           service_name=os.environ["ORACLE_SERVICE_NAME"])
    return oracledb.connect(user=os.environ["ORACLE_USER"], password=os.environ["ORACLE_PASSWORD"], dsn=dsn)


# --------------------------------------------------------------------------
# Utilidades de texto
# --------------------------------------------------------------------------
def limpo(s) -> str:
    if s is None:
        return ""
    return re.sub(r"\s+", " ", str(s).replace("\xa0", " ")).strip()


def sem_acento(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def bonito_depto(s: str) -> str:
    base = limpo(s).replace("_", " ")
    if sem_acento(base).lower() == "higiene beleza":
        return "Higiene & Beleza"
    return base


def bonito_categoria(s: str) -> str:
    base = limpo(s).replace("_", " ")
    if not base or base in ("0", "-"):
        return ""
    return base[:1].upper() + base[1:].lower()


def uf_do_supervisor(nome: str) -> str:
    m = re.match(r"^\s*(ES|BA|RJ)\b\s*[-–]", nome.upper())
    return m.group(1) if m else "ES"          # equipe sem prefixo é do ES (regra do BI original)


# --------------------------------------------------------------------------
# Venda (fato) — mensal x RCA x cliente x produto
# --------------------------------------------------------------------------
SQL_VENDA = f"""
SELECT TRUNC(m.dtmov,'MM')                                              AS mes,
       m.codusur, m.codcli, m.codprod,
       SUM(DECODE(m.codoper,'S',  m.qt*(m.punit+NVL(m.vloutros,0)), 0)) AS venda,
       SUM(DECODE(m.codoper,'ED', m.qt*(m.punit+NVL(m.vloutros,0)), 0)) AS devolucao,
       SUM(DECODE(m.codoper,'S',  m.qt, 0))                             AS qt_venda,
       SUM(DECODE(m.codoper,'ED', m.qt, 0))                             AS qt_dev,
       SUM(DECODE(m.codoper,'S', m.qt, -m.qt) * NVL(p.pesoliq,0))       AS peso,
       /* "com corte": só os dias até o dia de corte (o dia de ONTEM) de cada mês. Serve para comparar
          o mês em andamento com o MESMO período do ano/mês anterior, sem distorção. */
       SUM(CASE WHEN m.codoper = 'S' AND EXTRACT(DAY FROM m.dtmov) <= :dia
                THEN m.qt*(m.punit+NVL(m.vloutros,0)) ELSE 0 END)       AS venda_c,
       SUM(CASE WHEN m.codoper = 'ED' AND EXTRACT(DAY FROM m.dtmov) <= :dia
                THEN m.qt*(m.punit+NVL(m.vloutros,0)) ELSE 0 END)       AS devolucao_c,
       SUM(CASE WHEN m.codoper = 'S' AND EXTRACT(DAY FROM m.dtmov) <= :dia
                THEN m.qt ELSE 0 END)                                   AS qt_venda_c,
       SUM(CASE WHEN m.codoper = 'ED' AND EXTRACT(DAY FROM m.dtmov) <= :dia
                THEN m.qt ELSE 0 END)                                   AS qt_dev_c
  FROM pcmov m, pcprodut p
 WHERE m.codprod = p.codprod
   AND m.codoper IN ('S','ED')
   AND m.dtcancel IS NULL
   AND m.codfilial IN ({", ".join("'" + f + "'" for f in FILIAIS_VENDA)})
   AND m.dtmov >= :ini
   AND m.dtmov < TRUNC(SYSDATE)          /* só dias fechados: dados até ontem */
 GROUP BY TRUNC(m.dtmov,'MM'), m.codusur, m.codcli, m.codprod
 ORDER BY 1, 3, 4, 2          /* mês, cliente, produto: as linhas chegam agrupadas, o que acelera o motor */
"""


def extrair_venda(conn) -> dict:
    """Devolve colunas cruas (códigos do WinThor). Dinheiro em centavos inteiros."""
    cols = {"m": array("B"), "usur": array("I"), "cli": array("I"), "prod": array("I"),
            "v": array("i"), "d": array("i"), "qv": array("f"), "qd": array("f"), "w": array("f"),
            "vc": array("i"), "dc": array("i"), "qvc": array("f"), "qdc": array("f")}
    cur = conn.cursor()
    cur.arraysize = 50000
    t = time.time()
    ontem = date.today() - timedelta(days=1)
    cur.execute(SQL_VENDA, ini=datetime.combine(DATA_INICIAL, datetime.min.time()), dia=ontem.day)
    n = 0
    while True:
        lote = cur.fetchmany(50000)
        if not lote:
            break
        for mes, usur, cli, prod, venda, dev, qv, qd, peso, vc, dc, qvc, qdc in lote:
            cols["m"].append((mes.year - DATA_INICIAL.year) * 12 + mes.month - 1)
            cols["usur"].append(int(usur or 0))
            cols["cli"].append(int(cli or 0))
            cols["prod"].append(int(prod))
            cols["v"].append(int(round(float(venda or 0) * 100)))
            cols["d"].append(int(round(float(dev or 0) * 100)))
            cols["qv"].append(float(qv or 0))
            cols["qd"].append(float(qd or 0))
            cols["w"].append(float(peso or 0))
            cols["vc"].append(int(round(float(vc or 0) * 100)))
            cols["dc"].append(int(round(float(dc or 0) * 100)))
            cols["qvc"].append(float(qvc or 0))
            cols["qdc"].append(float(qdc or 0))
        n += len(lote)
        if n % 500000 < 50000:
            log(f"  venda: {n:,} linhas ({time.time() - t:.0f}s)".replace(",", "."))
    cur.close()
    log(f"venda extraída: {n:,} linhas em {time.time() - t:.0f}s".replace(",", "."))
    return cols


TIPOS_CACHE = {"m": "B", "usur": "I", "cli": "I", "prod": "I", "v": "i", "d": "i", "qv": "f", "qd": "f", "w": "f",
               "vc": "i", "dc": "i", "qvc": "f", "qdc": "f"}
CAMPOS_CACHE = tuple(TIPOS_CACHE)


def salvar_cache(cols: dict) -> None:
    with open(CACHE, "wb") as f:
        n = len(cols["m"])
        f.write(struct.pack("<I", n))
        for k in CAMPOS_CACHE:
            f.write(cols[k].tobytes())


def ler_cache() -> dict:
    tipos = TIPOS_CACHE
    cols = {}
    with open(CACHE, "rb") as f:
        n = struct.unpack("<I", f.read(4))[0]
        for k, tp in tipos.items():
            a = array(tp)
            a.frombytes(f.read(n * a.itemsize))
            cols[k] = a
    log(f"cache de venda lido: {n:,} linhas".replace(",", "."))
    return cols


# --------------------------------------------------------------------------
# Classificação do Trade (planilha mantida à mão) — vira dimensão de produto
# --------------------------------------------------------------------------
def achar_planilha_classificacao(relatorio: list | None = None) -> Path | None:
    """Primeira planilha acessível. `relatorio` (lista) recebe (caminho, motivo) de cada tentativa que falhou."""
    candidatas = list(ARQUIVOS_CLASSIFICACAO)
    if os.environ.get("TRADE_CLASSIFICACAO"):
        candidatas.append(Path(os.environ["TRADE_CLASSIFICACAO"]))
    for c in candidatas:
        try:
            if c.is_file():
                return c
            motivo = "não existe ou não está acessível"
        except OSError as e:
            motivo = f"{e.strerror or e} (erro {getattr(e, 'winerror', None) or e.errno})"
        if relatorio is not None:
            relatorio.append((str(c), motivo))
    return None


def ler_classificacao() -> dict:
    """codprod -> (departamento, categoria, subcategoria, marca, promocional); vazio se a planilha não está acessível"""
    import openpyxl
    arquivo = achar_planilha_classificacao()
    if arquivo is None:
        log("AVISO: planilha de classificação do Trade não encontrada em " + " nem em ".join(str(c) for c in ARQUIVOS_CLASSIFICACAO))
        return {}
    wb = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    out = {}
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0 or row is None or row[0] is None:
            continue
        try:
            cod = int(row[0])
        except (TypeError, ValueError):
            continue
        dep, cat, sub, marca, promo = (row[3], row[4], row[5], row[6], row[7])
        dep = limpo(dep)
        if dep in ("", "0"):
            dep = ""
        out[cod] = (dep, bonito_categoria(cat), bonito_categoria(sub), limpo(marca) if limpo(marca) not in ("0", "-") else "",
                    1 if limpo(promo).upper() == "P" else 0)
    log(f"classificação do Trade: {len(out):,} produtos".replace(",", "."))
    return out


def listar_imagens() -> dict:
    """codprod -> {arquivo, bytes, modificado}, direto da pasta FOTOS_ION; vazio se a pasta não está acessível"""
    import fotos
    pasta = fotos.localizar_pasta()
    if pasta is None:
        log("AVISO: pasta de fotos (FOTOS_ION) não encontrada em " + " nem em ".join(fotos.CANDIDATAS))
        return {}
    mapa, est = fotos.indexar(pasta)
    log(f"banco de imagens ({pasta}): {est['produtos']:,} produtos com foto, {est['repetidos']} códigos com mais de um arquivo, "
        f"{est['ignoradas']} entradas ignoradas".replace(",", "."))
    return mapa


# --------------------------------------------------------------------------
# Dimensões
# --------------------------------------------------------------------------
SQL_PRODUTOS = """
SELECT p.codprod, p.descricao, p.codauxiliar, p.codfornec, f.fornecedor, f.codcomprador, e.nome AS comprador,
       d.descricao AS depto_wint, p.pesoliq, p.embalagem, p.unidade, p.obs2, p.dtcadastro, p.dtexclusao, p.revenda,
       p.codepto, p.codsec, s.descricao AS secao_wint, cat.categoria AS categoria_wint, sub.subcategoria AS subcategoria_wint, m.marca AS marca_wint
  FROM pcprodut p, pcfornec f, pcempr e, pcdepto d, pcsecao s, pccategoria cat, pcsubcategoria sub, pcmarca m
 WHERE p.codfornec = f.codfornec (+)
   AND f.codcomprador = e.matricula (+)
   AND p.codepto = d.codepto (+)
   AND p.codsec = s.codsec (+)
   AND p.codsec = cat.codsec (+) AND p.codcategoria = cat.codcategoria (+)
   AND p.codsec = sub.codsec (+) AND p.codcategoria = sub.codcategoria (+) AND p.codsubcategoria = sub.codsubcategoria (+)
   AND p.codmarca = m.codmarca (+)
"""

SQL_CLIENTES = """
SELECT c.codcli, c.cliente, c.fantasia, c.codatv1, a.ramo, c.municent, c.estent, c.bloqueio, c.dtcadastro
  FROM pcclient c, pcativi a
 WHERE c.codatv1 = a.codativ (+)
"""

SQL_RCA = "SELECT u.codusur, u.nome, u.codsupervisor FROM pcusuari u"
SQL_SUPERVISOR = "SELECT s.codsupervisor, s.nome FROM pcsuperv s"


def consultar(conn, sql: str, **bind) -> list:
    cur = conn.cursor()
    cur.arraysize = 20000
    cur.execute(sql, bind)
    linhas = cur.fetchall()
    cur.close()
    return linhas


class Lista:
    """Dimensão por índice estável: valor -> posição."""
    def __init__(self):
        self.itens: list = []
        self.pos: dict = {}

    def id(self, chave, item=None):
        if chave not in self.pos:
            self.pos[chave] = len(self.itens)
            self.itens.append(item if item is not None else chave)
        return self.pos[chave]


# --------------------------------------------------------------------------
# Estoque x venda (base da Ruptura, Sem venda, Cobertura e Imagens)
# --------------------------------------------------------------------------
def _janela_semanal(i: int) -> str:
    if i == 0:
        faixa = "DTMOV >= TRUNC(SYSDATE,'WW') AND DTMOV < TRUNC(SYSDATE,'WW') + 7"
    else:
        ini = f"TRUNC(SYSDATE,'WW') - {7 * i}"
        fim = "TRUNC(SYSDATE,'WW')" if i == 1 else f"TRUNC(SYSDATE,'WW') - {7 * (i - 1)}"
        faixa = f"DTMOV >= {ini} AND DTMOV < {fim}"
    return f"""V{i} AS (SELECT CODPROD, SUM(QT) AS S FROM PCMOV
   WHERE CODOPER = 'S' AND CODFILIAL IN ('1','10','7') AND DTCANCEL IS NULL AND {faixa} GROUP BY CODPROD)"""


SQL_ESTOQUE = f"""
WITH EST AS (
  SELECT e.CODPROD,
         SUM(e.QTESTGER)                                   AS ESTOQUE,
         SUM(e.QTRESERV)                                   AS RESERVADO,
         SUM(e.QTBLOQUEADA - e.QTINDENIZ)                  AS BLOQ_SEMAV,
         SUM(e.QTINDENIZ)                                  AS AVARIA,
         SUM(e.QTESTGER - e.QTRESERV - e.QTBLOQUEADA)      AS DISPONIVEL,
         SUM((e.QTESTGER - e.QTRESERV - e.QTINDENIZ)
             * NVL(NULLIF(e.CUSTOFIN,0), NVL(e.CUSTOULTENT,0)))   AS VALOR_CUSTO
    FROM PCEST e
   WHERE e.CODFILIAL IN ({", ".join(str(f) for f in FILIAIS_ESTOQUE)})
   GROUP BY e.CODPROD
),
PRECO AS (SELECT CODPROD, PVENDA FROM PCTABPR WHERE NUMREGIAO = 1),
PEND AS (
  SELECT i.CODPROD, SUM(i.QTPEDIDA - i.QTENTREGUE) AS QTD_PENDENTE
    FROM PCPEDIDO p, PCITEM i
   WHERE p.NUMPED = i.NUMPED AND p.CODFILIAL IN ('1','10','7') AND (i.QTPEDIDA - i.QTENTREGUE) > 0
   GROUP BY i.CODPROD
),
{", ".join(_janela_semanal(i) for i in range(9))},
V56 AS (
  SELECT CODPROD, SUM(QT*(PUNIT+NVL(VLOUTROS,0))) AS VENDA_R
    FROM PCMOV
   WHERE CODOPER = 'S' AND DTCANCEL IS NULL AND CODFILIAL IN ({", ".join("'" + f + "'" for f in FILIAIS_VENDA)})
     AND DTMOV > TRUNC(SYSDATE) - {DIAS_VENDA_JANELA}
   GROUP BY CODPROD
),
ENT AS (
  SELECT CODPROD, MIN(DTMOV) AS PRIMEIRA, MAX(DTMOV) AS ULTIMA
    FROM PCMOV WHERE CODOPER = 'E' AND DTCANCEL IS NULL GROUP BY CODPROD
)
SELECT p.CODPROD,
       NVL(est.ESTOQUE,0), NVL(est.RESERVADO,0), NVL(est.BLOQ_SEMAV,0), NVL(est.AVARIA,0), NVL(est.DISPONIVEL,0),
       NVL(est.VALOR_CUSTO,0), pr.PVENDA, NVL(pe.QTD_PENDENTE,0),
       {", ".join(f"NVL(v{i}.S,0)" for i in range(9))},
       NVL(v56.VENDA_R,0), ent.PRIMEIRA, ent.ULTIMA
  FROM PCPRODUT p
  LEFT JOIN EST est  ON est.CODPROD = p.CODPROD
  LEFT JOIN PRECO pr ON pr.CODPROD = p.CODPROD
  LEFT JOIN PEND pe  ON pe.CODPROD = p.CODPROD
  {" ".join(f"LEFT JOIN V{i} v{i} ON v{i}.CODPROD = p.CODPROD" for i in range(9))}
  LEFT JOIN V56 v56  ON v56.CODPROD = p.CODPROD
  LEFT JOIN ENT ent  ON ent.CODPROD = p.CODPROD
 WHERE p.DTEXCLUSAO IS NULL AND p.REVENDA = 'S'
"""


def media_aparada(semanas: list[float]) -> float | None:
    """MEDIA do BI COMPRAS: descarta a melhor e a pior semana das 9."""
    tot = sum(semanas)
    if tot <= 1:
        return None
    com_venda = sum(1 for s in semanas if s > 0)
    if com_venda - 2 <= 0:
        return None
    return (tot - max(semanas) - min(semanas)) / (com_venda - 2)


# --------------------------------------------------------------------------
# Entradas (custo de compra) e conta corrente
# --------------------------------------------------------------------------
SQL_ENTRADAS = """
SELECT TRUNC(dtmov,'MM') AS mes, codprod,
       SUM(qt) AS qt,
       SUM(qt * (punit + NVL(vlfrete,0) + NVL(st,0))) AS valor
  FROM pcmov
 WHERE codoper = 'E' AND dtcancel IS NULL AND qt > 0 AND dtmov >= :ini
 GROUP BY TRUNC(dtmov,'MM'), codprod
"""

SQL_CC = """
SELECT c.numtranscrfor, c.data, c.codfornec, f.fornecedor, f.codcomprador, c.tipo, c.valor,
       c.historico1, c.numverba, u.nome AS usuario, c.dtestorno
  FROM pcmovcrfor c, pcfornec f, pcempr u
 WHERE c.codfornec = f.codfornec (+)
   AND c.codfunc = u.matricula (+)
"""

# Débito de "negociação" (regra do Power BI: texto do lançamento contém uma destas palavras;
# "ação" só vale se o texto NÃO começa com ALTERACAO).
PALAVRAS_NEGOCIACAO = ("REBAIXA", "VERBA", "VENCIMENTO", "PERDA", "COMPRA DE MERC.BONIFIC.", "NEGOC", "CAMPANHA")


def e_negociacao(hist: str) -> bool:
    t = sem_acento(hist or "").upper()
    if any(p in t for p in PALAVRAS_NEGOCIACAO):
        return True
    return ("ACAO" in t) and not t.startswith("ALTERACAO")


# --------------------------------------------------------------------------
# Leitura do Oracle (tudo menos a venda, que tem função própria)
# --------------------------------------------------------------------------
def ler_oracle(conn) -> dict:
    """Dimensões, base de estoque, entradas e conta corrente. Só SELECT."""
    log("dimensões…")
    o = {
        "sup": {int(r[0]): limpo(r[1]) for r in consultar(conn, SQL_SUPERVISOR)},
        "rca": {int(r[0]): (limpo(r[1]), int(r[2]) if r[2] is not None else None) for r in consultar(conn, SQL_RCA)},
        "cli": {int(r[0]): r for r in consultar(conn, SQL_CLIENTES)},
        "prod": {int(r[0]): r for r in consultar(conn, SQL_PRODUTOS)},
    }
    log("estoque x venda…")
    o["est"] = consultar(conn, SQL_ESTOQUE)
    log(f"  estoque: {len(o['est']):,} produtos".replace(",", "."))
    o["ent"] = consultar(conn, SQL_ENTRADAS, ini=datetime.combine(DATA_INICIAL, datetime.min.time()))
    log(f"  entradas: {len(o['ent']):,} produto x mês".replace(",", "."))
    o["cc"] = consultar(conn, SQL_CC)
    log(f"  conta corrente: {len(o['cc']):,} lançamentos".replace(",", "."))
    return o


# --------------------------------------------------------------------------
# O pacote na memória
# --------------------------------------------------------------------------
@dataclass
class Pacote:
    ref: date
    dims: dict
    est_cols: list
    est_out: list
    ent_out: list
    cc: dict                    # {"fornecedores", "cols_f", "lancamentos", "cols_l"}
    venda: list                 # colunas do fato: [(nome, tipo, array)]
    imagens: dict               # codprod -> nome do arquivo em FOTOS_ION
    hist_hoje: list             # linhas para historico_indicador (a foto de hoje)
    meta: dict
    n_venda: int = 0
    extras: dict = field(default_factory=dict)


def construir(cru: dict, o: dict, classif: dict, imagens: dict, ref: date | None = None, fontes: dict | None = None) -> Pacote:
    """Monta o pacote do painel a partir do que veio do Oracle (cru, o), da planilha do Trade (classif) e da pasta de fotos."""
    ref = ref or date.today()
    sup_rows, rca_rows, cli_rows, prod_rows = o["sup"], o["rca"], o["cli"], o["prod"]
    est_rows, ent_rows, cc_rows = o["est"], o["ent"], o["cc"]
    n_venda = len(cru["m"])

    # ---- índices -----------------------------------------------------
    UF = Lista()
    for u in ("ES", "BA", "RJ"):
        UF.id(u)
    SUP = Lista()       # supervisor/equipe
    RCA = Lista()
    CLI = Lista()
    RAMO = Lista()
    PROD = Lista()
    FORN = Lista()
    COMP = Lista()
    DEPTO = Lista()
    CAT = Lista()       # (depto, categoria)
    SUB = Lista()       # (categoria idx, subcategoria)
    MARCA = Lista()

    def idx_supervisor(cod: int | None) -> int:
        nome = sup_rows.get(cod) if cod is not None else None
        if nome is None:
            nome = "SEM EQUIPE"
        return SUP.id(cod if cod is not None else -1, [cod if cod is not None else -1, nome, UF.id(uf_do_supervisor(nome))])

    def idx_rca(cod: int) -> int:
        if cod in RCA.pos:
            return RCA.pos[cod]
        nome, sup = rca_rows.get(cod, (f"RCA {cod}", None))
        return RCA.id(cod, [cod, nome or f"RCA {cod}", idx_supervisor(sup)])

    def idx_cliente(cod: int) -> int:
        if cod in CLI.pos:
            return CLI.pos[cod]
        r = cli_rows.get(cod)
        if r:
            ramo = limpo(r[4]) or "Sem ramo"
            nome = limpo(r[2]) or limpo(r[1])
            nome = limpo(r[1]) or nome
            item = [cod, nome, RAMO.id(ramo), limpo(r[5]).title(), limpo(r[6]).upper()]
        else:
            item = [cod, f"Cliente {cod}", RAMO.id("Sem ramo"), "", ""]
        return CLI.id(cod, item)

    def idx_fornecedor(cod, nome, cod_comp, nome_comp) -> int:
        if cod is None:
            return -1
        cod = int(cod)
        if cod in FORN.pos:
            return FORN.pos[cod]
        if cod_comp is None or int(cod_comp) in COMPRADORES_EXCLUIDOS:
            ic = -1
        else:
            ic = COMP.id(int(cod_comp), [int(cod_comp), limpo(nome_comp)])
        return FORN.id(cod, [cod, limpo(nome), ic])

    SEM_CLASS = "Sem classificação"

    # Produtos novos ainda não estão na planilha do Trade. Para não ficarem num balde "sem classificação",
    # herdam o departamento do Trade da MAIORIA dos produtos já classificados da mesma seção do WinThor
    # (e, na falta, do mesmo departamento). Fica marcado em origem=1 e aparece no painel.
    from collections import Counter, defaultdict
    por_sec, por_dep = defaultdict(Counter), defaultdict(Counter)
    for cod, r in prod_rows.items():
        cl = classif.get(cod)
        if cl and cl[0]:
            por_sec[r[16]][bonito_depto(cl[0])] += 1
            por_dep[r[15]][bonito_depto(cl[0])] += 1
    mapa_sec = {k: c.most_common(1)[0][0] for k, c in por_sec.items() if c}
    mapa_dep = {k: c.most_common(1)[0][0] for k, c in por_dep.items() if c}

    def idx_produto(cod: int) -> int:
        if cod in PROD.pos:
            return PROD.pos[cod]
        r = prod_rows.get(cod)
        if not r:
            return PROD.id(cod, [cod, f"Produto {cod}", "", -1, DEPTO.id(SEM_CLASS), -1, -1, -1, 0, 0, 0.0, 2])
        (_, desc, ean, codf, fnome, codc, cnome, depto_w, peso, emb, unid, obs2, dtcad, dtexc, revenda,
         codepto, codsec, secao_w, cat_w, sub_w, marca_w) = r
        cl = classif.get(cod)
        origem = 0
        if cl and cl[0]:
            dep_nome = bonito_depto(cl[0])
            cat_nome, sub_nome, marca, promo = cl[1], cl[2], cl[3], cl[4]
        else:
            dep_nome = mapa_sec.get(codsec) or mapa_dep.get(codepto) or SEM_CLASS
            origem = 1 if dep_nome != SEM_CLASS else 2
            cat_nome = bonito_categoria(cat_w) or bonito_categoria(secao_w)
            sub_nome = bonito_categoria(sub_w)
            marca = limpo(marca_w).title()
            promo = cl[4] if cl else 0
        if not marca:
            marca = limpo(marca_w).title()
        idep = DEPTO.id(dep_nome)
        icat = CAT.id((idep, cat_nome or "Sem categoria"), [cat_nome or "Sem categoria", idep])
        isub = SUB.id((icat, sub_nome or "Sem subcategoria"), [sub_nome or "Sem subcategoria", icat])
        imar = MARCA.id(marca or "Sem marca")
        status = 1 if limpo(obs2).upper() == "FL" else 0          # 0 = ATIVO, 1 = FORA DE LINHA
        ean_s = str(int(ean)) if ean not in (None, "", 0) and str(ean).strip().isdigit() else limpo(ean)
        return PROD.id(cod, [cod, limpo(desc), ean_s, idx_fornecedor(codf, fnome, codc, cnome),
                             idep, icat, isub, imar, int(promo), status, float(peso or 0), origem])

    # produtos: ordem de aparição; vendas primeiro
    log("indexando venda…")
    m_arr, v_arr, d_arr = cru["m"], cru["v"], cru["d"]
    c_idx = array("I")
    p_idx = array("H")
    r_idx = array("H")
    cache_c, cache_p, cache_r = {}, {}, {}
    for i in range(n_venda):
        cod = cru["cli"][i]
        k = cache_c.get(cod)
        if k is None:
            k = cache_c[cod] = idx_cliente(cod)
        c_idx.append(k)
        cod = cru["prod"][i]
        k = cache_p.get(cod)
        if k is None:
            k = cache_p[cod] = idx_produto(cod)
        p_idx.append(k)
        cod = cru["usur"][i]
        k = cache_r.get(cod)
        if k is None:
            k = cache_r[cod] = idx_rca(cod)
        r_idx.append(k)
    log(f"  clientes {len(CLI.itens):,} · produtos {len(PROD.itens):,} · RCAs {len(RCA.itens):,}".replace(",", "."))

    # produtos da base de estoque que não venderam desde 2023 também entram
    for r in est_rows:
        idx_produto(int(r[0]))
    for r in ent_rows:
        idx_produto(int(r[1]))
    if len(PROD.itens) > 65535:
        raise SystemExit("mais de 65535 produtos: aumentar o tipo da coluna 'p'")

    # ---- estoque -----------------------------------------------------
    log("montando a base de estoque…")
    est_cols = ["p", "ativo", "promo", "disp", "bloq", "total", "avaria", "reserv", "preco", "valor_custo",
                "valor_pf", "venda56", "media", "cob", "sug", "pend", "prim", "ult", "img",
                "s0", "s1", "s2", "s3", "s4", "s5", "s6", "s7", "s8", "rup", "sv", "base"]
    est_out = []
    for r in est_rows:
        cod = int(r[0])
        p = PROD.pos[cod]
        pinfo = PROD.itens[p]
        (estoque, reserv, bloq_semav, avaria, disp, valor_custo, pvenda, pend) = (float(x or 0) for x in r[1:9])
        semanas = [float(x or 0) for x in r[9:18]]
        venda56 = float(r[18] or 0)
        prim, ult = r[19], r[20]
        ativo = 1 if pinfo[9] == 0 else 0
        promo = pinfo[8]
        total = disp + bloq_semav                         # "estoque total disponível" do original
        base = 1 if (ativo and not promo) else 0          # SKUs (p)
        recente = bool(ult and (ref - ult.date()).days < 365)
        rup = 1 if (base and recente and total <= 0) else 0
        primeira_ok = bool(prim and (ref - prim.date()).days > 14)
        sv = 1 if (ativo and venda56 <= 0 and primeira_ok and disp > 0) else 0
        media = media_aparada(semanas)
        cob = ((disp + bloq_semav) / media) if media else None
        sug = (media * 8 - disp) if media is not None and (media * 8 - disp) > 0 else None
        preco = float(pvenda or 0)
        est_out.append([p, ativo, promo, round(disp, 2), round(bloq_semav, 2), round(total, 2), round(avaria, 2),
                        round(reserv, 2), round(preco, 4), round(valor_custo, 2), round(total * preco / 1.35, 2),
                        round(venda56, 2), None if media is None else round(media, 3),
                        None if cob is None else round(cob, 2), None if sug is None else round(sug, 1),
                        round(pend, 1), iso(prim), iso(ult), 1 if cod in imagens else 0,
                        *[round(s, 1) for s in semanas], rup, sv, base])
    n_rup = sum(1 for e in est_out if e[est_cols.index("rup")])
    n_sv = sum(1 for e in est_out if e[est_cols.index("sv")])
    n_base = sum(1 for e in est_out if e[est_cols.index("base")])
    n_neg = sum(1 for e in est_out if e[est_cols.index("base")] and e[est_cols.index("total")] < 0)
    log(f"  base ativa não promocional: {n_base:,} · ruptura {n_rup} · sem venda {n_sv} · estoque negativo na base {n_neg}".replace(",", "."))

    # ---- entradas ----------------------------------------------------
    ent_out = []
    for mes, cod, qt, valor in ent_rows:
        mi = (mes.year - DATA_INICIAL.year) * 12 + mes.month - 1
        ent_out.append([PROD.pos[int(cod)], mi, round(float(qt), 3), round(float(valor), 2)])

    # ---- conta corrente ---------------------------------------------
    log("conta corrente…")
    ccf = {}           # fornecedor -> agregados
    usuarios = Lista()
    lanc = []
    # verba do último trimestre fechado (1º dia de [mês atual − 3] até o último dia do mês anterior)
    ini_tri = (date(ref.year, ref.month, 1) - timedelta(days=1)).replace(day=1)
    for _ in range(2):
        ini_tri = (ini_tri - timedelta(days=1)).replace(day=1)
    fim_tri = date(ref.year, ref.month, 1) - timedelta(days=1)
    for (trans, data, codf, fnome, codc, tipo, valor, hist, numverba, usuario, estorno) in cc_rows:
        if codf is None:
            continue       # estornos ENTRAM: o saldo corrido do relatório do WinThor (rotina 1825) inclui os dois lados
        codf = int(codf)
        valor = float(valor or 0)
        deb = valor if tipo == "D" else 0.0
        cre = valor if tipo == "C" else 0.0
        a = ccf.setdefault(codf, {"nome": limpo(fnome), "comp": codc, "deb": 0.0, "cre": 0.0, "trim": 0.0, "neg": 0.0})
        a["deb"] += deb
        a["cre"] += cre
        d0 = data.date()
        if deb and numverba is not None and ini_tri <= d0 <= fim_tri:
            a["trim"] += deb
        if deb and e_negociacao(hist) and ini_tri <= d0 <= fim_tri:
            a["neg"] += deb
        if d0 >= DATA_INICIAL_CC:
            lanc.append([d0.isoformat(), codf, 0 if tipo == "D" else 1, round(valor, 2), limpo(hist)[:140],
                         int(numverba) if numverba is not None else None, usuarios.id(limpo(usuario) or "—"),
                         int(trans), 1 if (tipo == "D" and e_negociacao(hist)) else 0, 1 if estorno is not None else 0])
    ativos_forn = {PROD.itens[e[0]][3] for e in est_out if e[1] == 1}
    ativos_cod = {FORN.itens[i][0] for i in ativos_forn if i >= 0}
    cc_forn = []
    for codf, a in sorted(ccf.items()):
        ic = -1
        if a["comp"] is not None and int(a["comp"]) in COMP.pos:
            ic = COMP.pos[int(a["comp"])]
        saldo = a["deb"] - a["cre"]
        verba_med = a["trim"] / 3.0
        cc_forn.append([codf, a["nome"], ic, round(a["deb"], 2), round(a["cre"], 2), round(saldo, 2),
                        round(verba_med, 2), round(a["neg"] / 3.0, 2), 1 if codf in ativos_cod else 0])
    log(f"  fornecedores com conta corrente: {len(cc_forn)} · lançamentos exibidos: {len(lanc):,}".replace(",", "."))
    cc = {"fornecedores": cc_forn,
          "cols_f": ["cod", "nome", "comprador", "debito", "credito", "saldo", "verba_media", "negociacao_media", "ativo"],
          "lancamentos": lanc,
          "cols_l": ["data", "forn", "tipo", "valor", "historico", "verba", "usuario", "transacao", "negociacao", "estorno"]}

    # ---- foto de hoje dos indicadores (vira linha de historico_indicador) -------------
    por_comp: dict = {}
    for e in est_out:
        pinfo = PROD.itens[e[0]]
        forn = FORN.itens[pinfo[3]] if pinfo[3] >= 0 else None
        ic = forn[2] if forn else -1
        d = por_comp.setdefault(ic, {"skus": 0, "rup": 0, "sv": 0, "ativos": 0, "parado": 0.0, "parado_pf": 0.0})
        d["ativos"] += e[1]
        d["skus"] += e[est_cols.index("base")]
        d["rup"] += e[est_cols.index("rup")]
        d["sv"] += e[est_cols.index("sv")]
        if e[est_cols.index("sv")]:
            d["parado"] += e[est_cols.index("valor_custo")]
            d["parado_pf"] += e[est_cols.index("valor_pf")]
    hist_hoje = []
    for ic, d in sorted(por_comp.items()):
        cod_c, nome_c = (-1, "Sem comprador") if ic < 0 else (COMP.itens[ic][0], COMP.itens[ic][1])
        hist_hoje.append((ref, cod_c, nome_c, d["ativos"], d["skus"], d["rup"], d["sv"], round(d["parado"], 2), round(d["parado_pf"], 2)))

    # ---- dimensões e meta ------------------------------------------------
    ontem = ref - timedelta(days=1)
    meses = [f"{DATA_INICIAL.year + (i // 12)}-{(i % 12) + 1:02d}" for i in range(max(m_arr) + 1)]
    dims = {
        "ref": ref.isoformat(),
        "corte": {"ate": ontem.isoformat(), "dia": ontem.day, "mes": ontem.month,
                  "mi": (ontem.year - DATA_INICIAL.year) * 12 + ontem.month - 1},
        "meses": meses,
        "uf": UF.itens,
        "equipes": SUP.itens,                           # [cod, nome, uf]
        "rcas": RCA.itens,                              # [cod, nome, equipe]
        "clientes": CLI.itens,                          # [cod, nome, ramo, municipio, uf]
        "ramos": RAMO.itens,
        "deptos": DEPTO.itens,
        "categorias": CAT.itens,                        # [nome, depto]
        "subcategorias": SUB.itens,                     # [nome, categoria]
        "marcas": MARCA.itens,
        "fornecedores": FORN.itens,                     # [cod, nome, comprador]
        "compradores": COMP.itens,                      # [cod, nome]
        "produtos": PROD.itens,                         # [cod, desc, ean, forn, depto, cat, sub, marca, promo, status, peso, origem(0 trade|1 winthor|2 nenhum)]
        "usuarios_cc": usuarios.itens,
    }
    tot_ano: dict = {}
    for i in range(n_venda):
        ano = DATA_INICIAL.year + m_arr[i] // 12
        a = tot_ano.setdefault(ano, [0, 0])
        a[0] += v_arr[i]
        a[1] += d_arr[i]
    meta = {
        "gerado_em": datetime.now().isoformat(timespec="seconds"),
        "ref": ref.isoformat(),
        "periodo": {"ini": DATA_INICIAL.isoformat(), "fim": ref.isoformat()},
        "linhas_venda": n_venda,
        "filiais_venda": list(FILIAIS_VENDA),
        "contagens": {"produtos": len(PROD.itens), "clientes": len(CLI.itens), "rcas": len(RCA.itens),
                      "fornecedores": len(FORN.itens), "compradores": len(COMP.itens), "estoque": len(est_out),
                      "ruptura": n_rup, "sem_venda": n_sv, "base_ativa": n_base, "entradas": len(ent_out),
                      "cc_fornecedores": len(cc_forn), "cc_lancamentos": len(lanc),
                      "imagens": len(imagens), "estoque_com_imagem": sum(e[est_cols.index("img")] for e in est_out)},
        "venda_liquida_por_ano": {str(a): round((v - d) / 100, 2) for a, (v, d) in sorted(tot_ano.items())},
        "classificados_pct": round(100.0 * sum(1 for p in PROD.itens if DEPTO.itens[p[4]] != SEM_CLASS) / max(1, len(PROD.itens)), 1),
        "fontes": fontes or {},
    }
    log("venda líquida por ano: " + " · ".join(f"{a}: {v:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")
                                               for a, v in meta["venda_liquida_por_ano"].items()))
    venda = [("m", "u1", m_arr), ("c", "u4", c_idx), ("p", "u2", p_idx), ("r", "u2", r_idx),
             ("v", "i4", v_arr), ("d", "i4", d_arr), ("qv", "f4", cru["qv"]), ("qd", "f4", cru["qd"]), ("w", "f4", cru["w"]),
             ("vc", "i4", cru["vc"]), ("dc", "i4", cru["dc"]), ("qvc", "f4", cru["qvc"]), ("qdc", "f4", cru["qdc"])]
    return Pacote(ref=ref, dims=dims, est_cols=est_cols, est_out=est_out, ent_out=ent_out, cc=cc, venda=venda,
                  imagens={int(c): v["arquivo"] for c, v in imagens.items()}, hist_hoje=hist_hoje, meta=meta, n_venda=n_venda)


# --------------------------------------------------------------------------
# Serialização do pacote
# --------------------------------------------------------------------------
def iso(d) -> str | None:
    if d is None:
        return None
    if isinstance(d, str):
        return d
    if isinstance(d, datetime):
        return d.date().isoformat()
    return d.isoformat()


def json_bytes(obj) -> bytes:
    return json.dumps(obj, ensure_ascii=False, separators=(",", ":")).encode("utf-8")


def colunar_gz(colunas: list[tuple[str, str, array]]) -> bytes:
    """Formato: [uint32 tamanho do cabeçalho][cabeçalho JSON][pad até 8][colunas alinhadas em 8], tudo gzip."""
    import io
    n = len(colunas[0][2])
    cab = {"n": n, "cols": []}
    off = 0
    blocos = []
    for nome_c, tipo, arr in colunas:
        dados = arr.tobytes()
        cab["cols"].append({"nome": nome_c, "tipo": tipo, "off": off})
        pad = (-len(dados)) % 8
        blocos.append((dados, pad))
        off += len(dados) + pad
    cab_b = json.dumps(cab, separators=(",", ":")).encode()
    topo = struct.pack("<I", len(cab_b)) + cab_b
    topo += b"\0" * ((-len(topo)) % 8)
    saida = io.BytesIO()
    with gzip.GzipFile(fileobj=saida, mode="wb", compresslevel=6, mtime=0) as f:
        f.write(topo)
        for dados, pad in blocos:
            f.write(dados)
            if pad:
                f.write(b"\0" * pad)
    return saida.getvalue()


def historico_json(linhas) -> list:
    """Linhas de historico_indicador -> o que o painel lê: [{data, comprador: {codcomprador: {...}}}]."""
    por_data: dict = {}
    for d, cod, nome, ativos, skus, rup, sv, parado, parado_pf in linhas:
        dia = por_data.setdefault(iso(d), {})
        dia[str(int(cod))] = {"nome": nome or "", "ativos": int(ativos), "skus": int(skus), "rup": int(rup), "sv": int(sv),
                              "parado": round(float(parado), 2), "parado_pf": round(float(parado_pf), 2)}
    return [{"data": d, "comprador": por_data[d]} for d in sorted(por_data)]


def serializar(p: Pacote, historico: list) -> dict[str, bytes]:
    """Nome do arquivo -> conteúdo. JSON sem compressão; venda.bin.gz já comprimido."""
    blobs = {
        "dims.json": json_bytes(p.dims),
        "estoque.json": json_bytes({"cols": p.est_cols, "linhas": p.est_out}),
        "entradas.json": json_bytes({"cols": ["p", "m", "qt", "valor"], "linhas": p.ent_out}),
        "cc.json": json_bytes(p.cc),
        "historico.json": json_bytes(historico),
        "imagens.json": json_bytes({str(c): a for c, a in sorted(p.imagens.items())}),
        "venda.bin.gz": colunar_gz(p.venda),
    }
    meta = dict(p.meta)
    meta["bytes"] = {k: len(v) for k, v in blobs.items()}
    blobs["meta.json"] = json_bytes(meta)
    return blobs


def gravar_arquivos(blobs: dict[str, bytes], saida: Path = SAIDA) -> None:
    """Modo sem banco: grava o pacote em ../web/dados/ (o servidor local lê de lá quando não alcança o banco)."""
    saida.mkdir(parents=True, exist_ok=True)
    for nome, dados in blobs.items():
        (saida / nome).write_bytes(dados)


# --------------------------------------------------------------------------
def main() -> int:
    """Modo sem banco: Oracle -> arquivos em ../web/dados/. A carga de verdade é o sync_trade.py."""
    usar_cache = "--usar-cache" in sys.argv
    t0 = time.time()
    carregar_env()
    conn = conectar()
    log("conectado ao Oracle")
    if usar_cache and CACHE.is_file():
        cru = ler_cache()
    else:
        cru = extrair_venda(conn)
        salvar_cache(cru)
    o = ler_oracle(conn)
    conn.close()
    classif = ler_classificacao()
    imagens = listar_imagens()
    p = construir(cru, o, classif, imagens, fontes={"modo": "arquivos"})
    # histórico local (formato novo: por código de comprador); sem banco, cada carga acrescenta a foto de hoje
    antigo = []
    cam = SAIDA / "historico.json"
    if cam.is_file():
        try:
            antigo = [h for h in json.loads(cam.read_text(encoding="utf-8")) if h.get("data") != p.ref.isoformat()
                      and all(k.lstrip("-").isdigit() and isinstance(v, dict) and "nome" in v for k, v in h.get("comprador", {}).items())]
        except Exception:
            antigo = []
    linhas = [(h["data"], int(k), v["nome"], v["ativos"], v["skus"], v["rup"], v["sv"], v["parado"], v["parado_pf"])
              for h in antigo for k, v in h["comprador"].items()] + list(p.hist_hoje)
    blobs = serializar(p, historico_json(linhas))
    gravar_arquivos(blobs)
    log(f"pronto em {time.time() - t0:.0f}s · {sum(len(b) for b in blobs.values()) / 1e6:.1f} MB em {SAIDA}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
